import sys, os, subprocess, re, shutil, math
import numpy as np
from PIL import Image
from docx import Document
from docx.shared import Pt, Mm
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
sys.path.insert(0,'/home/claude/build')
TMP='/tmp/claude-0/pg/'; os.makedirs(TMP,exist_ok=True)
TOPM=BOTM=Mm(12)
def topdf(path):
    subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',os.path.dirname(path),path],capture_output=True)
    return path[:-5]+'.pdf'
def render(pdf):
    for f in os.listdir(TMP):
        if f.startswith('r-'): os.remove(TMP+f)
    subprocess.run(['pdftoppm','-r','72','-gray','-png',pdf,TMP+'r'])
    fs=sorted(f for f in os.listdir(TMP) if f.startswith('r-'))
    return [np.array(Image.open(TMP+f)) for f in fs]
def bottom(a):
    rows=np.where((a<200).sum(1)>0)[0]
    return int(rows.max()) if len(rows) else 0
def pagetext(pdf,n): return re.sub(r'\W+','',subprocess.run(['pdftotext','-f',str(n),'-l',str(n),pdf,'-'],capture_output=True,text=True).stdout.lower())
def body(d): return [e for e in d.element.body.iterchildren() if e.tag!=qn('w:sectPr')]
TASKRE=re.compile(r'^\s*(task|part|activity|exercise|handout|worksheet)\s*[0-9a-z]?\b',re.I)
def is_head(d,e):
    if e.tag!=qn('w:p'): return False
    p=Paragraph(e,d)
    if p.style.name.startswith('Heading') or p.style.name=='Title': return True
    rs=[r for r in p.runs if r.text.strip()]
    return bool(rs) and bool(rs[0].bold) and bool(TASKRE.match(p.text)) and len(p.text)<160
def forced(d,e):
    return e.tag==qn('w:p') and 'student' in Paragraph(e,d).text.lower() and Paragraph(e,d).style.name.startswith('Heading')
def blocks(d):
    ks=body(d); starts=[]
    for i,e in enumerate(ks):
        if is_head(d,e) and not (i>0 and is_head(d,ks[i-1])): starts.append(i)
    def titleish(e): return e.tag==qn('w:p') and Paragraph(e,d).style.name in ('Heading 1','Title')
    starts=[s for s in starts if s>0 and not titleish(ks[s])]
    # first block = title + intro + first task
    starts=starts[1:] if len(starts)>1 else starts
    if not starts: return ks,[(0,len(ks))]
    bl=[(0,starts[0])]+[(a,b) for a,b in zip(starts,starts[1:]+[len(ks)])]
    return ks,bl
def prep(d):
    for s in d.sections:
        s.left_margin=s.right_margin=Mm(13); s.top_margin=TOPM; s.bottom_margin=BOTM
        s.header_distance=s.footer_distance=Mm(6)
    allp=list(d.paragraphs)+[p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]
    for p in allp:
        pf=p.paragraph_format; h=p.style.name.startswith('Heading') or p.style.name=='Title'
        pf.space_before=Pt(6 if h else 0); pf.space_after=Pt(2 if h else 3); pf.line_spacing=1.0
        pf.keep_with_next=None; pf.page_break_before=None
        for r in p.runs:
            if r.font.size is not None and r.font.size<Pt(11): r.font.size=Pt(11)
    # remove explicit page breaks
    for br in d.element.body.iter(qn('w:br')):
        if br.get(qn('w:type'))=='page': br.getparent().remove(br)
def first_para(e):
    return e if e.tag==qn('w:p') else next(e.iter(qn('w:p')))
def set_break(d,e,on=True): Paragraph(first_para(e),d).paragraph_format.page_break_before=on
def keep_block(d,ks,a,b):
    for e in ks[a:b-1]:
        for p in ([e] if e.tag==qn('w:p') else e.iter(qn('w:p'))): Paragraph(p,d).paragraph_format.keep_with_next=True
def measure(f):
    d=Document(f); prep(d); ks,bl=blocks(d)
    for a,b in bl[1:]: set_break(d,ks[a])
    out=TMP+'m.docx'; d.save(out); pdf=topdf(out); imgs=render(pdf)
    H=imgs[0].shape[0]; top=int(TOPM.pt); cap=H-top-int(BOTM.pt)
    # map blocks to pages by heading text
    heads=[re.sub(r'\W+','',Paragraph(first_para(ks[a]),d).text.lower())[:25] for a,b in bl]
    texts=[pagetext(pdf,i+1) for i in range(len(imgs))]
    startpg=[]; pg=0
    for i,h in enumerate(heads):
        if i==0: startpg.append(0); continue
        while pg<len(texts) and not (h and texts[pg].startswith(h[:15])) : pg+=1
        if pg>=len(texts): return None
        startpg.append(pg)
    hs=[]
    for i,s in enumerate(startpg):
        e=(startpg[i+1]-1) if i+1<len(startpg) else len(imgs)-1
        hs.append((e-s)*cap + bottom(imgs[e])-top + 4)
    return hs,cap
def plan(hs,cap,forcedset=frozenset()):
    n=len(hs); INF=float('inf'); gap=0
    def fits(i,j): # blocks i..j-1 on one page
        return sum(hs[i:j])<=cap*0.97
    # min pages via DP, oversize blocks alone
    best={n:(0,0.0,None)}
    for i in range(n-1,-1,-1):
        cand=[]
        for j in range(i+1,n+1):
            if j>i+1 and (not fits(i,j) or (j-1) in forcedset): break
            fill=sum(hs[i:j]); pages=max(1,math.ceil(fill/cap)) if j==i+1 else 1
            used=fill-(pages-1)*cap
            p2,cost2,_=best[j]
            cand.append((p2+pages, cost2+(cap-used)**2, j))
        best[i]=min(cand)
    groups=[]; i=0
    while i<n: j=best[i][2]; groups.append((i,j)); i=j
    # a side that would be mostly empty (<35%) is worse than letting a task run over: merge it with the next side
    k=0
    while k<len(groups)-1:
        a,b=groups[k]
        if sum(hs[a:b])<cap*0.35 and (b) not in forcedset:
            groups[k]=(a,groups[k+1][1]); del groups[k+1]
        else: k+=1
    return best[0][0],groups
def apply(f,groups,hs,cap):
    d=Document(f); prep(d); ks,bl=blocks(d)
    for gi,(i,j) in enumerate(groups):
        if gi>0: set_break(d,ks[bl[i][0]])
        for k in range(i,j):
            a,b=bl[k]; 
            if hs[k]<=cap*0.97 and sum(hs[i:j])<=cap*0.97: keep_block(d,ks,a,b)
        # spread leftover space between tasks on this page
        fill=sum(hs[i:j]); left=cap*0.97-fill
        if j-i>1 and left>0 and fill<=cap:
            add=min(36,left/(j-i))
            for k in range(i+1,j):
                p=Paragraph(first_para(ks[bl[k][0]]),d); p.paragraph_format.space_before=Pt(6+add)
    d.save(f)
f=sys.argv[1]; base=f[:-5]
before=int(re.search(r'Pages:\s+(\d+)',subprocess.run(['pdfinfo',base+'.pdf'],capture_output=True,text=True).stdout)[1])
r=measure(f)
if r is None: print(os.path.basename(f),'could not map blocks; skipped'); sys.exit()
hs,cap=r
_d=Document(f); prep(_d); _ks,_bl=blocks(_d)
fs=frozenset(k for k,(a,b) in enumerate(_bl) if k>0 and forced(_d,_ks[a]))
pages,groups=plan(hs,cap,fs)
fills=[round(sum(hs[i:j])/cap*100) for i,j in groups]
over=[k for k,h in enumerate(hs) if h>cap*0.97]
shutil.copy(f,TMP+'backup.docx')
flow=max(1,math.ceil(sum(hs)/(cap*0.97)))
mode='no splits'
if len(groups)>flow and sum(hs)/(len(groups)*cap)<0.65:
    # keeping every task whole wastes too much paper: let tasks flow, only forced breaks
    groups=[]; st=0
    for k in range(1,len(hs)+1):
        if k==len(hs) or k in fs: groups.append((st,k)); st=k
    mode='flowing (keeping tasks whole would waste a side)'
    d=Document(f); prep(d); ks,bl=blocks(d)
    for gi,(i,j) in enumerate(groups):
        if gi>0: set_break(d,ks[bl[i][0]])
    d.save(f)
else:
    apply(f,groups,hs,cap)
pdf=topdf(f)
after=int(re.search(r'Pages:\s+(\d+)',subprocess.run(['pdfinfo',pdf],capture_output=True,text=True).stdout)[1])
print(f"{f}: {before} -> {after} sides, {mode}; fill % per side {fills}; tasks longer than a side: {len(over)}")

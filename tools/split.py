import sys, copy, os, subprocess
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
W=qn
def kids(d): return [e for e in d.element.body.iterchildren() if e.tag!=qn('w:sectPr')]
def style(d,e):
    return Paragraph(e,d).style.name if e.tag==qn('w:p') else ''
def text(d,e): return Paragraph(e,d).text if e.tag==qn('w:p') else ''
def ans_index(d):
    ks=kids(d)
    for i,e in enumerate(ks):
        if style(d,e).startswith('Heading') and any(k in text(d,e).lower() for k in ['answer key','answers (teacher','teacher answer']):
            return i
    return None
def keep_tasks(d):
    ks=kids(d)
    for i,e in enumerate(ks):
        nxt=ks[i+1] if i+1<len(ks) else None
        if nxt is None: break
        nst=style(d,nxt)
        if nst.startswith('Heading'): continue  # task ends here
        targets=[e] if e.tag==qn('w:p') else e.iter(qn('w:p'))
        for p in targets:
            Paragraph(p,d).paragraph_format.keep_with_next=True
        if e.tag==qn('w:tbl'):
            for tr in e.iter(qn('w:tr')):
                trPr=tr.find(qn('w:trPr'))
                if trPr is None: trPr=OxmlElement('w:trPr'); tr.insert(0,trPr)
                if trPr.find(qn('w:cantSplit')) is None: trPr.append(OxmlElement('w:cantSplit'))
def pages(path):
    out=os.path.dirname(path); subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',out,path],capture_output=True)
    pdf=path[:-5]+'.pdf'
    return int([l for l in subprocess.run(['pdfinfo',pdf],capture_output=True,text=True).stdout.splitlines() if l.startswith('Pages')][0].split()[1])
f=sys.argv[1]; base=f[:-5]
d=Document(f); i=ans_index(d)
if i is not None:
    a=Document(f); ka=kids(a)
    h1=next(j for j,e in enumerate(ka) if style(a,e).startswith('Heading 1'))
    for e in ka[h1+1:i]: e.getparent().remove(e)
    p=Paragraph(ka[h1],a); p.add_run(' – Answers')
    a.save(base+'-answers.docx'); print('answers pages',pages(base+'-answers.docx'))
    for e in kids(d)[i:]: e.getparent().remove(e)
d.save(f); p0=pages(f)
k=Document(f); keep_tasks(k); k.save(base+'.keep.docx'); p1=pages(base+'.keep.docx')
if p1<=p0:
    os.replace(base+'.keep.docx',f); os.replace(base+'.keep.pdf',base+'.pdf')
else:
    os.remove(base+'.keep.docx'); os.remove(base+'.keep.pdf')
print(f,'pages',p0,'->',p1, 'kept' if p1<=p0 else 'reverted (would add pages)')

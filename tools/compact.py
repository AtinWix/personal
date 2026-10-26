import sys, os, subprocess, re, shutil
from docx import Document
from docx.shared import Pt, Mm
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def pdf(path):
    subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',os.path.dirname(path) or '.',path],capture_output=True)
    return path[:-5]+'.pdf'
def pagetexts(p):
    n=int(re.search(r'Pages:\s+(\d+)',subprocess.run(['pdfinfo',p],capture_output=True,text=True).stdout)[1])
    return [subprocess.run(['pdftotext','-f',str(i),'-l',str(i),p,'-'],capture_output=True,text=True).stdout for i in range(1,n+1)]
def norm(s): return re.sub(r'\W+','',s.lower())
def headings(d,levels):
    return [p.text for p in d.paragraphs if p.style.name in levels and p.text.strip()]
def splits(d,pdfp,levels):
    pages=[norm(t) for t in pagetexts(pdfp)]
    hs=headings(d,levels); pos=[]
    for h in hs:
        k=norm(h)[:30]; pg=next((i for i,t in enumerate(pages) if k in t),None); pos.append((h,pg))
    bad=[]
    for (h,p),(h2,p2) in zip(pos,pos[1:]+[(None,len(pages)-1)]):
        if p is None or p2 is None: continue
        if h2 is None: end=p2
        else:
            idx=pages[p2].find(norm(h2)[:30]); end=p2 if idx>2 else p2-1  # next heading near top => task ended on previous page
        if end>p: bad.append(h)
    return len(pages),bad
def task_levels(d):
    names={p.style.name for p in d.paragraphs}
    return {'Heading 3'} if 'Heading 3' in names else {'Heading 2'}

def compact(d):
    for s in d.sections:
        s.left_margin=s.right_margin=Mm(13); s.top_margin=s.bottom_margin=Mm(12)
        s.header_distance=s.footer_distance=Mm(6)
    for p in d.paragraphs+[p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
        pf=p.paragraph_format
        h=p.style.name.startswith('Heading') or p.style.name=='Title'
        pf.space_before=Pt(6 if h else 0); pf.space_after=Pt(2 if h else 3); pf.line_spacing=1.0
        for r in p.runs:
            if r.font.size is not None and r.font.size<Pt(11): r.font.size=Pt(11)
    for st in d.styles:
        try:
            if st.font.size is not None and st.font.size<Pt(11) and st.type==1: st.font.size=Pt(11)
        except Exception: pass
def keep(d,levels,only=None):
    body=d.element.body; ks=[e for e in body.iterchildren() if e.tag!=qn('w:sectPr')]
    def isbound(e):
        if e.tag!=qn('w:p'): return False
        n=Paragraph(e,d).style.name
        return n.startswith('Heading')
    cur=None
    for i,e in enumerate(ks[:-1]):
        if e.tag==qn('w:p') and Paragraph(e,d).style.name in levels: cur=Paragraph(e,d).text
        if isbound(ks[i+1]): continue
        if only is not None and cur not in only: continue
        ps=[e] if e.tag==qn('w:p') else list(e.iter(qn('w:p')))
        for p in ps: Paragraph(p,d).paragraph_format.keep_with_next=True
        if e.tag==qn('w:tbl'):
            for tr in e.iter(qn('w:tr')):
                trPr=tr.find(qn('w:trPr'))
                if trPr is None: trPr=OxmlElement('w:trPr'); tr.insert(0,trPr)
                if trPr.find(qn('w:cantSplit')) is None: trPr.append(OxmlElement('w:cantSplit'))

f=sys.argv[1]; base=f[:-5]; tmp='/tmp/claude-0/cmp/'; os.makedirs(tmp,exist_ok=True)
d0=Document(f); lv=task_levels(d0)
n0,b0=splits(d0,base+'.pdf',lv)
def build(do_compact,only):
    d=Document(f)
    if do_compact: compact(d)
    if only: keep(d,lv,only)
    out=tmp+os.path.basename(base)+'.try.docx'; d.save(out); p=pdf(out)
    n,b=splits(d,p,lv); return n,b,out,p
best=(n0,len(b0),None,None,b0,'orig')
n,b,out,p=build(True,set())
if (n,len(b))<best[:2]:
    best=(n,len(b),tmp+'best.docx',tmp+'best.pdf',b,'compact'); shutil.copy(out,best[2]); shutil.copy(p,best[3])
base_compact = best[5]=='compact'
fixed=set()
for t in list(best[4]):
    trial=fixed|{t}
    n,b,out,p=build(base_compact,trial)
    if n<=best[0] and len(b)<best[1]:
        fixed=trial; best=(n,len(b),tmp+'best.docx',tmp+'best.pdf',b,('compact+' if base_compact else '')+'keep'); shutil.copy(out,best[2]); shutil.copy(p,best[3])
if best[5]!='orig':
    shutil.copy(best[2],f); shutil.copy(best[3],base+'.pdf')
print(f"{os.path.basename(f)}: {n0} sides, {len(b0)} split -> {best[5]}: {best[0]} sides, {best[1]} split {best[4][:4]}")
